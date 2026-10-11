"""Shared fixtures: a small multi-page sample document plus PDF/DOCX builders."""
import io

import pytest
from docx import Document

from src.document_loader import fitz  # same PyMuPDF module the app uses

SAMPLE_TEXTS = [
    "Python is a high-level programming language. Python was created by Guido van Rossum and first "
    "released in 1991. Its design emphasizes code readability and simple syntax.",
    "The Eiffel Tower is a wrought-iron lattice tower in Paris, France. It was designed by the "
    "engineering company of Gustave Eiffel and completed in 1889 for the World's Fair.",
    "Photosynthesis is the process by which green plants convert sunlight, water and carbon dioxide "
    "into glucose and oxygen. Chlorophyll inside the chloroplasts absorbs the light energy.",
    "BM25 is a ranking function used by search engines to score documents against a query. It relies "
    "on term frequency, inverse document frequency and document length normalization.",
    "The Moon orbits the Earth roughly every twenty-seven days. Its gravity causes the ocean tides, "
    "and its surface is covered with craters formed by ancient impacts.",
    "The Nile is one of the longest rivers in Africa. It flows northward through Egypt into the "
    "Mediterranean Sea, and its annual floods once made the valley fertile.",
]


@pytest.fixture
def sample_pages():
    return [{"page": i + 1, "section": None, "text": text} for i, text in enumerate(SAMPLE_TEXTS)]


@pytest.fixture
def sample_document(sample_pages):
    return {"name": "sample.pdf", "extension": ".pdf", "total_pages": 6, "pages": sample_pages}


@pytest.fixture
def make_pdf():
    """make_pdf(["page 1 text", "", "page 3 text"]) -> PDF bytes (empty string = blank page)."""

    def _make(page_texts):
        pdf = fitz.open()
        for text in page_texts:
            page = pdf.new_page()
            if text:
                page.insert_text((72, 72), text, fontsize=10)
        data = pdf.tobytes()
        pdf.close()
        return data

    return _make


@pytest.fixture
def make_docx():
    """make_docx([("Heading", "Paragraph"), (None, "Paragraph")]) -> DOCX bytes."""

    def _make(items):
        document = Document()
        for heading, paragraph in items:
            if heading:
                document.add_heading(heading, level=1)
            if paragraph:
                document.add_paragraph(paragraph)
        buffer = io.BytesIO()
        document.save(buffer)
        return buffer.getvalue()

    return _make
